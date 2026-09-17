from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# Create presentation
prs = Presentation()

# Define Colors
BG_COLOR = RGBColor(240, 235, 220)       # Beige background
DEEP_GREEN = RGBColor(20, 65, 45)        # Deep green for shapes/cards
LIGHT_TEXT = RGBColor(245, 245, 240)     # Off-white for text inside cards
DARK_TEXT = RGBColor(30, 30, 30)         # Dark for text on beige background

def set_font(run, color, size, is_title=False):
    run.font.name = 'Calibri'
    run.font.color.rgb = color
    run.font.size = Pt(size)
    if is_title:
        run.font.bold = True

def apply_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR

# Define slide contents
slides_data = [
    {
        "type": "title",
        "title": "Clase Introductoria",
        "subtitle": "Primer acercamiento a los contenidos del taller"
    },
    {
        "type": "content",
        "title": "Preparación Inicial",
        "body": "• Iniciar sesión en la cuenta de invitado (siempre limpia)\n• Acceso a cuenta de Google de Tumo\n• Acceso a cuenta de tumo.world\n\n¡Estamos listos para empezar!"
    },
    {
        "type": "content",
        "title": "Herramientas de Trabajo",
        "body": "• Google Classroom\n• Google Drive\n• VSCode\n• Tumo Path\n\nClaves en la vida diaria del taller"
    },
    {
        "type": "content",
        "title": "Google Classroom y Drive",
        "body": "Google Classroom:\n• Visualizar materiales, entregar ejercicios, recibir feedback y consultas\n\nGoogle Drive:\n• Repositorio personal\n• Importante: guardar antes de finalizar\n  (las cuentas borran archivos al cerrar sesión)"
    },
    {
        "type": "content",
        "title": "VS Code y Tumo Path",
        "body": "VS Code:\n• Editor de código para escribir y desarrollar proyectos\n• Configurar \"Autosave\" para evitar pérdidas de información\n\nTumo Path:\n• Abrir para marcar el presente cada día"
    },
    {
        "type": "content",
        "title": "Pasos para el Inicio",
        "body": "1. Iniciar sesión como invitado en la iMac\n2. tumo.world -> Establecer estado \"Presente\"\n3. Google Drive -> Descargar/subir archivos de proyectos\n4. VSCode -> Establecer \"Auto Save\""
    },
    {
        "type": "content",
        "title": "Pasos para el Cierre",
        "body": "1. Guardar el proyecto correctamente\n2. Subir el proyecto a Drive\n3. Subir tareas entregables a Classroom\n4. Cerrar sesión de todas las cuentas e iMac"
    },
    {
        "type": "content",
        "title": "Importante",
        "body": "• Dediquen tiempo y esfuerzo para lograr buenos proyectos\n• Disfruten la experiencia y aprovechen para aprender\n• No importa cuántas veces fallen, sino qué aprenden de ello\n\n\"Nada que valga la pena sale perfecto a la primera...\""
    },
    {
        "type": "title",
        "title": "¡ÉXITOS!",
        "subtitle": ""
    }
]

# Create slides
for slide_data in slides_data:
    if slide_data["type"] == "title":
        slide_layout = prs.slide_layouts[6] # Blank slide to build from scratch
        slide = prs.slides.add_slide(slide_layout)
        apply_background(slide)
        
        # Add a deep green card in the center
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(1.5), Inches(2), Inches(7), Inches(3.5)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = DEEP_GREEN
        card.line.fill.background() # No border
        
        # Add Title Text Box over the card
        txBox = slide.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(7), Inches(1.5))
        tf = txBox.text_frame
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = slide_data["title"]
        set_font(run, LIGHT_TEXT, 48, True)
        
        # Add Subtitle Text Box over the card
        if slide_data["subtitle"]:
            txBox2 = slide.shapes.add_textbox(Inches(2), Inches(3.5), Inches(6), Inches(1.5))
            tf2 = txBox2.text_frame
            tf2.word_wrap = True
            p2 = tf2.paragraphs[0]
            p2.alignment = PP_ALIGN.CENTER
            run2 = p2.add_run()
            run2.text = slide_data["subtitle"]
            set_font(run2, LIGHT_TEXT, 22)
            
    elif slide_data["type"] == "content":
        slide_layout = prs.slide_layouts[6] # Blank slide
        slide = prs.slides.add_slide(slide_layout)
        apply_background(slide)
        
        # Add a decorative deep green stripe on the left
        stripe = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0), Inches(0), Inches(0.4), Inches(7.5)
        )
        stripe.fill.solid()
        stripe.fill.fore_color.rgb = DEEP_GREEN
        stripe.line.fill.background()
        
        # Add Title on the beige background
        txBox_title = slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(8), Inches(1))
        tf_title = txBox_title.text_frame
        p_title = tf_title.paragraphs[0]
        run_title = p_title.add_run()
        run_title.text = slide_data["title"]
        set_font(run_title, DEEP_GREEN, 36, True)
        
        # Add a deep green card for the content body
        card = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, 
            Inches(1), Inches(1.6), Inches(8), Inches(5.2)
        )
        card.fill.solid()
        card.fill.fore_color.rgb = DEEP_GREEN
        card.line.fill.background()
        
        # Add Content text box inside the card
        txBox_body = slide.shapes.add_textbox(Inches(1.3), Inches(1.9), Inches(7.4), Inches(4.6))
        tf_body = txBox_body.text_frame
        tf_body.word_wrap = True
        
        lines = slide_data["body"].split('\n')
        for i, line in enumerate(lines):
            if i == 0:
                p_body = tf_body.paragraphs[0]
            else:
                p_body = tf_body.add_paragraph()
            
            run_body = p_body.add_run()
            run_body.text = line
            set_font(run_body, LIGHT_TEXT, 24)

prs.save('Introduccion_Taller_Redesign.pptx')
print("Presentation created successfully as Introduccion_Taller_Redesign.pptx")
